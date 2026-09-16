"""Build soil-samples.json from the lab exports instead of from an old page.

    python dev/jd_rx_soil_import.py "C:/Users/acces/.grain-map-secrets/soil-import"

The folder is the unzipped "Studer Farm Soil Samples (2014-2026)" archive.
Three sources go in, one file comes out:

  TonySents_SoilSampleData_2014_to_2025_Exported_Fieldalytics/*.shp+.dbf
      624 grid sampling events, 2014-2025, one shapefile per field per
      date, WGS84 points with the full Waypoint / Midwest panel. This is
      the same data dev/jd_rx_soil.py scraped out of the chat-built page,
      but from the lab's own export, with Na and base saturation that the
      page had dropped.

  Farmers Edge Labs - Spring 2026 1-Acre Grids/*.xlsx
      2026 one-acre grids, Mehlich-3, with latitude/longitude per bag.

  2026 Soil Data - Claude Transcribed Excel Version/ALL 2026 - Waypoint
  Soil Sample Results.xlsx
      71 WHOLE-FIELD composites - one row of averages per field, no
      coordinates. 53 of them are South org fields, which had no soil
      data at all. These are written with an empty point list and
      kind="whole"; dev/jd_build_rx.py spreads each one over its
      boundary so the rate still varies with yield history, and the page
      labels the field as a composite rather than a grid.

Writes ~/.grain-map-secrets/soil-samples.json (private, outside the repo).

Two things worth knowing about the numbers:

  * A lab value of exactly 0 means "not in this test package", not zero
    ppm - the 2014 events report 0 for every micronutrient. Zeros are
    dropped, so a spot that was never tested for zinc stays blank
    instead of reading as deficient.
  * Farmers Edge reports Sikora buffer pH; the 2014-2025 Waypoint events
    report SMP. ISU treats them as interchangeable for the lime tables,
    so both land in "bph".

The whole-field rows name their field in 14 characters, and some of those
names have drifted from Operations Center ("Richrdsn31Glen", "Beilenbrg70Ben"),
so they are matched to Deere fields by farm name plus the acreage in the
name, never by string equality. Whatever does not match is listed at the
end and stored under "unmatched" - nothing is dropped silently.
"""
from __future__ import annotations

import datetime as dt
import difflib
import json
import pathlib
import re
import struct
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from jd_common import SECRETS, now_iso, read_json, write_private  # noqa: E402

OUTPUT = SECRETS / "soil-samples.json"
OPS = SECRETS / "rx-ops.json"
# Hand-written {"lab name": "Operations Center field name"} for rows the
# matcher cannot place. Anything listed here wins over the guessing below.
ALIASES = SECRETS / "soil-field-aliases.json"
FLEET = SECRETS / "fleet.json"

SHP_DIR = "TonySents_SoilSampleData_2014_to_2025_Exported_Fieldalytics"
FEL_DIR = "Farmers Edge Labs - Spring 2026 1-Acre Grids"
WP_2026 = "2026 Soil Data - Claude Transcribed Excel Version"

ORG_OF = {"Studer Farms-North": "5294", "Studer Farms-South": "567678",
          "Studer Farms - North": "5294", "Studer Farms - South": "567678"}

# Lab column -> the short key the prescription page uses.
NUT = {
    "P": "p", "K": "k", "PH": "ph", "BUFFERPH": "bph", "OM": "om",
    "CEC": "cec", "CA": "ca", "MG": "mg", "S": "s", "ZN": "zn", "MN": "mn",
    "FE": "fe", "CU": "cu", "B": "b", "NA": "na", "AL": "al", "CL": "cl",
    "NO3-N": "no3n", "NH4-N": "nh4n", "%CLAY": "clay", "%SAND": "sand",
    "%SILT": "silt", "BS": "bs", "BS_CA": "bs_ca", "BS_K": "bs_k",
    "BS_MG": "bs_mg", "BS_NA": "bs_na", "BS_H": "bs_h",
    "P1": "p1", "P2": "p2", "PMEHCOL": "pmeh", "POLSEN": "polsen",
    "CA_AA": "ca_aa", "MG_AA": "mg_aa", "S_AA": "s_aa", "ZN_DTPA": "zn_dtpa",
}

# Waypoint and Midwest Labs report the same nutrient under different
# column names, and the page reads one key. First column that has a value
# wins. P is the one that matters: "P" (Waypoint Bray-1), "P1" (Midwest
# Bray-1) and "PMEHCOL" (Mehlich-3) all read on the same ISU scale, so
# they can share a key. Olsen P does NOT - it runs about half of Bray on
# the same soil and has its own interpretation table - but no event here
# reports Olsen without Bray alongside it, so "polsen" is carried as an
# extra and never used as "p".
SAME_AS = {
    "p": ("P", "P1", "PMEHCOL"),
    "ca": ("CA", "CA_AA"),
    "mg": ("MG", "MG_AA"),
    "s": ("S", "S_AA"),
    "zn": ("ZN", "ZN_DTPA"),
}


def num(v):
    """A lab number, or None when the value is missing or a 0 placeholder."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if not v or v in ("None", "NA", "N/A", "-"):
            return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f == 0.0:
        return None
    return round(f, 4)


# --------------------------------------------------------------------- dbf/shp

def read_dbf(path: pathlib.Path) -> list[dict]:
    b = path.read_bytes()
    nrec, hlen, rlen = struct.unpack("<IHH", b[4:12])
    cols, off = [], 32
    while b[off] != 0x0D:
        f = b[off:off + 32]
        cols.append((f[:11].split(b"\x00")[0].decode("latin-1"), chr(f[11]), f[16]))
        off += 32
    rows = []
    for r in range(nrec):
        rec = b[hlen + r * rlen: hlen + (r + 1) * rlen]
        if not rec or rec[:1] == b"*":       # deleted record
            continue
        o, row = 1, {}
        for name, _typ, ln in cols:
            row[name] = rec[o:o + ln].decode("latin-1").strip()
            o += ln
        rows.append(row)
    return rows


def read_shp_points(path: pathlib.Path) -> list[tuple[float, float]]:
    """(lon, lat) per record. These exports are all shape type 1, points."""
    b = path.read_bytes()
    pts, o = [], 100
    while o + 8 <= len(b):
        _num, wlen = struct.unpack(">2i", b[o:o + 8])
        body = b[o + 8: o + 8 + wlen * 2]
        if struct.unpack("<i", body[:4])[0] == 1:
            pts.append(struct.unpack("<2d", body[4:20]))
        else:
            pts.append((None, None))
        o += 8 + wlen * 2
    return pts


def from_shapefiles(root: pathlib.Path, fields: dict) -> tuple[int, int]:
    n_sets = n_pts = 0
    for shp in sorted((root / SHP_DIR).glob("*.shp")):
        m = re.match(r"(.+?)_(.+)_SampleData_(\d{4}-\d\d-\d\d)\.shp", shp.name)
        if not m:
            print(f"  skipped (name): {shp.name}")
            continue
        county, field, date = m.group(1), m.group(2), m.group(3)
        rows = read_dbf(shp.with_suffix(".dbf"))
        coords = read_shp_points(shp)
        if len(rows) != len(coords):
            print(f"  {shp.name}: {len(rows)} records but {len(coords)} points")
        pts, lab = [], None
        for row, (x, y) in zip(rows, coords):
            if x is None:
                continue
            lab = lab or (row.get("LAB") if row.get("LAB") != "None" else None)
            pt = {"x": round(x, 7), "y": round(y, 7),
                  "id": num(row.get("ID_NUM")) or row.get("ID")}
            for col, key in NUT.items():
                if col in row:
                    v = num(row[col])
                    if v is not None:
                        pt[key] = v
            for key, cols in SAME_AS.items():
                if pt.get(key) is None:
                    for col in cols:
                        v = num(row.get(col))
                        if v is not None:
                            pt[key] = v
                            break
            pts.append(pt)
        if not pts:
            continue
        rec = fields.setdefault(field, {"county": county, "sets": []})
        rec.setdefault("county", county)
        rec["sets"].append({"d": date, "n": len(pts), "lab": lab,
                            "kind": "grid", "avg": averages(pts), "pts": pts})
        n_sets += 1
        n_pts += len(pts)
    return n_sets, n_pts


def averages(pts: list[dict]) -> dict:
    """Field average per nutrient over the points that reported it."""
    tot: dict = {}
    for p in pts:
        for k, v in p.items():
            if k in ("x", "y", "id") or not isinstance(v, (int, float)):
                continue
            s, n = tot.get(k, (0.0, 0))
            tot[k] = (s + v, n + 1)
    return {k: round(s / n, 2) for k, (s, n) in tot.items() if n}


# ------------------------------------------------------------------ 2026 grids

FEL_COLS = {
    "P M3 ppm": "p", "K M3 ppm": "k", "% OM": "om", "ph 1:1": "ph",
    "sikora BpH": "bph", "SO4 M3 ppm": "s", "CA M3 ppm": "ca",
    "Mg M3 ppm": "mg", "Na M3 ppm": "na", "Cu M3 ppm": "cu",
    "Fe M3 ppm": "fe", "Mn M3 ppm": "mn", "Zn M3 ppm": "zn",
    "B M3 ppm": "b", "Al M3 ppm": "al", "CEC mg/110g": "cec",
    "% Base Sat": "bs", "% K": "bs_k", "% Ca": "bs_ca", "% Mg": "bs_mg",
    "% Na": "bs_na", "EC ds/m": "ec",
}


def from_farmers_edge(root: pathlib.Path, fields: dict,
                      byorg: dict, aliases: dict) -> tuple[int, int, list]:
    """Only one of the three 2026 Farmers Edge workbooks carries Latitude
    and Longitude. The other two are the same one-acre grid pulled without
    coordinates, so their bags cannot be placed; they come in as a
    whole-field composite of the bags rather than being thrown away."""
    import openpyxl
    n_grid = n_pts = 0
    whole = []
    for xl in sorted((root / FEL_DIR).glob("*.xlsx")):
        wb = openpyxl.load_workbook(xl, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            head = [str(c).strip() if c is not None else "" for c in rows[0]]
            idx = {h: i for i, h in enumerate(head)}
            if "Field" not in idx:
                continue
            located = "Latitude" in idx and "Longitude" in idx
            pts, field, date = [], None, None
            for r in rows[1:]:
                if r[idx["Field"]] is None:
                    continue
                field = field or str(r[idx["Field"]]).strip()
                d = r[idx["Delivery Date"]] if "Delivery Date" in idx else None
                date = date or (d.date().isoformat() if hasattr(d, "date")
                                else str(d)[:10])
                pt = {"id": r[idx["Bag ID"]] if "Bag ID" in idx else None}
                if located:
                    if r[idx["Latitude"]] is None:
                        continue
                    pt["x"] = round(float(r[idx["Longitude"]]), 7)
                    pt["y"] = round(float(r[idx["Latitude"]]), 7)
                for col, key in FEL_COLS.items():
                    if col in idx:
                        v = num(r[idx[col]])
                        if v is not None:
                            pt[key] = v
                pts.append(pt)
            if not pts or not field:
                continue
            avg = averages(pts)
            if located:
                rec = fields.setdefault(field, {"county": None, "sets": []})
                rec["sets"].append({"d": date, "n": len(pts),
                                    "lab": "FarmersEdgeLabs", "kind": "grid",
                                    "avg": avg, "pts": pts})
                n_grid += 1
                n_pts += len(pts)
                continue
            name, how = resolve(field, byorg, aliases)
            if not name:
                whole.append({"name": field, "grower": "Farmers Edge 2026",
                              "report": xl.name, "d": date, "why": how,
                              "avg": avg})
                continue
            rec = fields.setdefault(name, {"county": None, "sets": []})
            rec["sets"].append({
                "d": date, "n": len(pts), "lab": "FarmersEdgeLabs",
                "kind": "whole", "avg": avg, "pts": [], "src": field,
                "report": "1-acre grid, no coordinates in the export",
                "match": how, "field_id": field_id(byorg, name),
                "org": org_of(byorg, name)})
    return n_grid, n_pts, whole


def field_id(byorg: dict, name: str) -> str | None:
    for names in byorg.values():
        if name in names:
            return names[name]
    return None


def org_of(byorg: dict, name: str) -> str | None:
    for org, names in byorg.items():
        if name in names:
            return org
    return None


def resolve(name: str, byorg: dict, aliases: dict,
            org: str | None = None) -> tuple[str | None, str]:
    """Match a lab field name against Deere, in one org or in all of them."""
    if org:
        return match_field(name, byorg.get(org, {}), aliases)
    seen: dict = {}
    for names in byorg.values():
        seen.update(names)
    return match_field(name, seen, aliases)


# ------------------------------------------------------- 2026 whole-field rows

WP_COLS = {
    "OM%": "om", "ENR": "enr", "CEC_meq_100g": "cec", "pH": "ph",
    "Buffer_pH": "bph", "P_ppm": "p", "K_ppm": "k", "Ca_ppm": "ca",
    "Mg_ppm": "mg", "S_ppm": "s", "B_ppm": "b", "Cu_ppm": "cu",
    "Fe_ppm": "fe", "Mn_ppm": "mn", "Zn_ppm": "zn", "Na_ppm": "na",
    "pctK": "bs_k", "pctCa": "bs_ca", "pctMg": "bs_mg", "pctH": "bs_h",
    "pctNa": "bs_na",
}


def report_date(report_no: str) -> str | None:
    """Waypoint report numbers are yy-ddd-nnnn, the day of the year the
    samples came in: 25-336-0501 is 2 December 2025."""
    m = re.match(r"(\d\d)-(\d{1,3})-", str(report_no).strip())
    if not m:
        return None
    year, day = 2000 + int(m.group(1)), int(m.group(2))
    if not 1 <= day <= 366:
        return None
    return (dt.date(year, 1, 1) + dt.timedelta(days=day - 1)).isoformat()


def deere_names() -> dict:
    """{org: {name: field_id}} from the operations pull."""
    if not OPS.exists():
        print(f"  {OPS.name} not found - whole-field rows cannot be matched")
        return {}
    ops = json.load(open(OPS, encoding="utf-8"))
    out: dict = {}
    for fid, f in ops.get("fields", {}).items():
        out.setdefault(f.get("org"), {})[f.get("name")] = fid
    return out


def split_name(name: str) -> tuple[str, str]:
    """"Richrdsn264Cha" -> ("richrdsn", "264"). The farm, then the acres."""
    m = re.match(r"([A-Za-z_]+)(\d+)", name.strip())
    return (m.group(1).lower(), m.group(2)) if m else (name.strip().lower(), "")


def load_aliases() -> dict:
    if not ALIASES.exists():
        return {}
    out = json.load(open(ALIASES, encoding="utf-8"))
    return {k.strip(): v.strip() for k, v in out.items()
            if v and not k.startswith("_")}


def match_field(name: str, candidates: dict, aliases: dict) -> tuple[str | None, str]:
    """Deere field name for a whole-field row, and how it was matched.

    Exact prefix first (the row name is truncated at 14 characters), then
    same acreage with a farm name close enough to be the same farm -
    which is what catches Beilenbrg/Bielenbrg and Wheatley75Cumm ->
    Wheatley75/EffingtonCummins1."""
    if name in aliases:
        alias = aliases[name]
        return (alias, "alias") if alias in candidates else (None, "alias not a field")
    hit = [n for n in candidates if n.startswith(name)]
    if len(hit) == 1:
        return hit[0], "prefix"
    farm, acres = split_name(name)
    if not acres:
        return (None, "ambiguous prefix" if hit else "no match")
    near = []
    for n in candidates:
        f2, a2 = split_name(n)
        if a2 != acres:
            continue
        if f2.startswith(farm) or farm.startswith(f2) or \
                difflib.SequenceMatcher(None, farm, f2).ratio() >= 0.75:
            near.append(n)
    if len(near) == 1:
        return near[0], "farm+acres"
    if len(hit) > 1:
        return None, "ambiguous prefix"
    return None, "ambiguous acres" if near else "no match"


REPORTS = SECRETS / "soil-reports.json"


def load_reports() -> dict:
    """{report number: what the PDF says about that field}.

    dev/jd_rx_soil_pdf.py reads the lab's own report PDFs. Their map pages
    carry the field's FULL name, its county, its acres and its centroid -
    everything the spreadsheet had to cut down to 14 characters. With the
    centroid a row can be placed on the map instead of guessed at by
    name."""
    if not REPORTS.exists():
        return {}
    data = json.load(open(REPORTS, encoding="utf-8"))
    fields = data.get("fields", {})
    out = {}
    for rep, r in data.get("reports", {}).items():
        f = fields.get(r.get("field_id") or "") or {}
        out[rep] = {"field": f.get("field") or r.get("field"),
                    "county": f.get("county"), "acres": f.get("acres"),
                    "lat": f.get("lat"), "lon": f.get("lon"),
                    "samples": len(r.get("samples") or [])}
    return out


def deere_boundaries() -> list:
    """[(name, org, geometry)] for every field with a boundary."""
    fleet = read_json(FLEET, None)
    if not fleet:
        return []
    sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
    from jd_build_rx import field_geometry
    out = []
    for f in fleet.get("fields", []):
        if not f.get("rings"):
            continue
        g = field_geometry(f)
        if g:
            out.append((f["name"], str(f.get("org")), g))
    return out


def _pip(x: float, y: float, ring: list) -> bool:
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i][0], ring[i][1]
        xj, yj = ring[j][0], ring[j][1]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def _contains(geom: dict, x: float, y: float) -> bool:
    polys = [geom["coordinates"]] if geom["type"] == "Polygon" else geom["coordinates"]
    for rings in polys:
        if _pip(x, y, rings[0]) and not any(_pip(x, y, h) for h in rings[1:]):
            return True
    return False


def match_by_centroid(lon: float, lat: float, boundaries: list,
                      org: str | None) -> tuple[str | None, str]:
    """The Deere field the report's centroid falls in.

    A centroid can land in a neighbour when a field wraps around one, so
    a hit is only taken when exactly one field contains the point."""
    hits = [n for n, o, g in boundaries
            if (not org or o == org) and _contains(g, lon, lat)]
    if len(hits) == 1:
        return hits[0], "centroid"
    if len(hits) > 1:
        return None, f"centroid in {len(hits)} fields"
    return None, "centroid outside every boundary"


def from_waypoint_2026(root: pathlib.Path, fields: dict, byorg: dict,
                       aliases: dict, reports: dict,
                       boundaries: list) -> tuple[int, list]:
    """The 2026 whole-field rows, placed on Deere fields.

    Each row is tried four ways, best evidence first: a hand-written
    alias, the lab report's CENTROID falling inside one boundary, the
    report's full field name, and last the spreadsheet's own name, which
    the lab cut to 14 characters."""
    import openpyxl
    unmatched, n_sets = [], 0
    for xl in sorted((root / WP_2026).glob("*.xlsx")):
        wb = openpyxl.load_workbook(xl, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            head = [str(c).strip() if c is not None else "" for c in rows[0]]
            idx = {h: i for i, h in enumerate(head)}
            if "FieldName" not in idx:
                continue
            for r in rows[1:]:
                raw = r[idx["FieldName"]]
                if raw is None:
                    continue
                raw = str(raw).strip()
                grower = str(r[idx["Grower"]]).strip() if "Grower" in idx else ""
                org = ORG_OF.get(grower)
                avg = {}
                for col, key in WP_COLS.items():
                    if col in idx:
                        v = num(r[idx[col]])
                        if v is not None:
                            avg[key] = v
                if not avg:
                    continue
                report = str(r[idx["Report_No"]]).strip() if "Report_No" in idx else ""
                date = report_date(report) or "2026-01-01"
                pdf = reports.get(report) or {}
                full = pdf.get("field")

                name, how = (None, "")
                if raw in aliases:
                    name, how = match_field(raw, byorg.get(org, {}), aliases)
                if not name and boundaries and pdf.get("lat") and pdf.get("lon"):
                    name, how = match_by_centroid(pdf["lon"], pdf["lat"],
                                                  boundaries, org)
                if not name and full:
                    name, how2 = match_field(full, byorg.get(org, {}), {})
                    how = how2 + " (report name)" if name else how or how2
                if not name:
                    name, how2 = match_field(raw, byorg.get(org, {}), {})
                    how = how2 if name else how or how2
                if not name:
                    unmatched.append({"name": raw, "full": full,
                                      "grower": grower, "report": report,
                                      "d": date, "why": how, "avg": avg,
                                      "county": pdf.get("county"),
                                      "acres": pdf.get("acres"),
                                      "lat": pdf.get("lat"), "lon": pdf.get("lon")})
                    continue
                rec = fields.setdefault(name, {"county": pdf.get("county"),
                                               "sets": []})
                if pdf.get("county"):
                    rec["county"] = rec.get("county") or pdf["county"]
                rec["sets"].append({
                    "d": date, "n": pdf.get("samples"),
                    "lab": "WaypointAnalyticalIowa",
                    "kind": "whole", "avg": avg, "pts": [],
                    "src": full or raw, "report": report, "match": how,
                    "field_id": byorg.get(org, {}).get(name), "org": org})
                n_sets += 1
        wb.close()
    return n_sets, unmatched


def drop_duplicate_wholes(fields: dict) -> int:
    """The 18 North rows in the 2026 Waypoint sheet are the field averages
    of November grid events that came in from Fieldalytics with all their
    points. Keep the points, drop the average."""
    n = 0
    for rec in fields.values():
        grids = [dt.date.fromisoformat(s["d"]) for s in rec["sets"]
                 if s.get("kind") == "grid"]
        keep = []
        for s in rec["sets"]:
            if s.get("kind") == "whole":
                d = dt.date.fromisoformat(s["d"])
                if any(abs((d - g).days) <= 14 for g in grids):
                    n += 1
                    continue
            keep.append(s)
        rec["sets"] = keep
    return n


# ------------------------------------------------------------------------ main

def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    root = pathlib.Path(sys.argv[1]).expanduser()
    if not root.is_dir():
        raise SystemExit(f"not a folder: {root}")

    fields: dict = {}
    print(f"reading {root}")
    s1, p1 = from_shapefiles(root, fields)
    print(f"  Fieldalytics shapefiles: {s1} events, {p1:,} points")
    byorg, aliases = deere_names(), load_aliases()
    if aliases:
        print(f"  {len(aliases)} name aliases from {ALIASES.name}")
    s2, p2, fel_whole = from_farmers_edge(root, fields, byorg, aliases)
    print(f"  Farmers Edge 2026 grids: {s2} located events, {p2:,} points")
    reports = load_reports()
    boundaries = deere_boundaries() if reports else []
    if reports:
        print(f"  {len(reports)} lab reports read from the PDFs "
              f"(dev/jd_rx_soil_pdf.py), {len(boundaries)} boundaries to place them in")
    s3, unmatched = from_waypoint_2026(root, fields, byorg, aliases,
                                       reports, boundaries)
    print(f"  Waypoint 2026 whole-field: {s3} matched, {len(unmatched)} unmatched")
    unmatched += fel_whole

    dropped = drop_duplicate_wholes(fields)
    if dropped:
        print(f"  dropped {dropped} whole-field rows that are the field "
              f"average of a grid event already imported")

    for rec in fields.values():
        rec["sets"].sort(key=lambda s: s["d"])

    write_private(OUTPUT, {"generated_at": now_iso(), "source": root.name,
                           "fields": fields, "unmatched": unmatched},
                  separators=(",", ":"))
    n_sets = sum(len(f["sets"]) for f in fields.values())
    n_pts = sum(len(s["pts"]) for f in fields.values() for s in f["sets"])
    print(f"wrote {OUTPUT}: {len(fields)} fields, {n_sets} sampling events, "
          f"{n_pts:,} grid points")
    if unmatched:
        print("\nwhole-field rows with no Operations Center field:")
        for u in unmatched:
            extra = ""
            if u.get("full") and u["full"] != u["name"]:
                extra = f"  full name {u['full']}"
            if u.get("lat"):
                extra += f"  at {u['lat']:.5f},{u['lon']:.5f}"
                if u.get("acres"):
                    extra += f"  {u['acres']:.1f} ac"
            print(f"  {u['name']:16s} {u['grower']:20s} {u['d']}  "
                  f"{u['why']}{extra}")


if __name__ == "__main__":
    main()
